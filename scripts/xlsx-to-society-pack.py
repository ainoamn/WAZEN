#!/usr/bin/env python3
"""Convert a Khalfan-style society workbook into a Wazen importSocietyPack JSON.

Requires: pip install openpyxl
Usage: python scripts/xlsx-to-society-pack.py "نظام جمعية أبناء خلفان.xlsx" -o pack.json
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path


def iso(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    return text[:10] if text else "2026-04-01"


def is_opening(name: str, typ: str, note: str, amount: float) -> bool:
    blob = f"{name}|{typ}|{note}"
    if re.search(r"مبالغ\s*سابقة|الجمعية السابقة|رصيد افتتاح", blob, re.I):
        return True
    return amount >= 1000 and bool(re.search(r"سابق|مرشد", blob))


def convert(path: Path, space_name: str) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    members = []
    for i, row in enumerate(wb["المساهمون"].iter_rows(values_only=True)):
        if i == 0 or not row[1]:
            continue
        name = str(row[1]).strip()
        status = str(row[4] or "").strip()
        if status and status != "نشط":
            continue
        phone = str(row[2]).strip() if row[2] is not None else ""
        email = str(row[3]).strip() if row[3] else ""
        item = {"displayName": name, "joinedAt": iso(row[6])}
        if phone:
            item["phone"] = phone
        if email:
            item["email"] = email
        members.append(item)

    payments = []
    opening = None
    for i, row in enumerate(wb["سجل الإيرادات"].iter_rows(values_only=True)):
        if i == 0 or not row[2]:
            continue
        name = str(row[2]).strip()
        typ = str(row[3] or "").strip()
        note = str(row[9] or "").strip()
        amount = float(row[4] or 0)
        when = iso(row[1])
        if is_opening(name, typ, note, amount):
            if opening is None or float(opening["amount"]) < amount:
                opening = {
                    "amount": amount,
                    "date": when,
                    "note": "مبالغ أخرى — رصيد افتتاحي"
                    + (f" · {note}" if note else (f" · {typ}" if typ else "")),
                }
            continue
        pay = {"date": when, "memberName": name, "amount": amount}
        if note or typ:
            pay["note"] = note or typ
        payments.append(pay)

    expenses = []
    holdings = []
    for i, row in enumerate(wb["المصروفات"].iter_rows(values_only=True)):
        if i == 0 or row[2] is None:
            continue
        typ = str(row[3] or "").strip()
        details = str(row[4] or "").strip()
        blob = f"{typ} {details}"
        if re.search(r"وديع|استثمار|أسهم", blob, re.I):
            holdings.append({
                "date": iso(row[1]),
                "amount": float(row[2]),
                "name": details or typ or "وديعة",
                "kind": "deposit",
                "note": typ or None,
            })
            continue
        item = {
            "date": iso(row[1]),
            "amount": float(row[2]),
            "description": details or typ or "مصروف",
        }
        beneficiary = str(row[5]).strip() if row[5] else ""
        if beneficiary:
            item["beneficiary"] = beneficiary
        expenses.append(item)

    return {
        "spaceName": space_name,
        "type": "society",
        "currency": "OMR",
        "monthlyContribution": 2,
        "durationMonths": 12,
        "startsAt": "2026-04-01",
        "dueDay": 1,
        "openingBalance": opening,
        "members": members,
        "payments": payments,
        "expenses": expenses,
        "holdings": holdings,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("-o", "--output", type=Path, required=True)
    parser.add_argument("--name", default="جمعية أبناء خلفان")
    args = parser.parse_args()
    pack = convert(args.xlsx, args.name)
    args.output.write_text(json.dumps(pack, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"wrote {args.output} · members={len(pack['members'])} "
        f"payments={len(pack['payments'])} expenses={len(pack['expenses'])} "
        f"holdings={len(pack.get('holdings') or [])} "
        f"opening={pack['openingBalance']}"
    )


if __name__ == "__main__":
    main()
